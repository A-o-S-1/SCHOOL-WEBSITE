"""Tests for result access and staff-only result-management tools."""
import datetime

from django.test import Client, TestCase, override_settings
from django.urls import reverse

from apps.academics.models import Department, Subject
from apps.results.models import (
    AcademicSession,
    ClassRoom,
    Student,
    Term,
)


@override_settings(STORAGES={
    "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
    "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
})
class ITStaffDashboardTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        from django.contrib.auth.models import Group, User
        from apps.results.models import ResultEntry, TermResult

        cls.it_group = Group.objects.get(name="IT Staff")
        cls.it_user = User.objects.create_user(
            username="it-operator", password="long-test-password"
        )
        cls.it_user.groups.add(cls.it_group)
        cls.unauthorized_user = User.objects.create_user(
            username="ordinary-user", password="long-test-password"
        )
        cls.inactive_it_user = User.objects.create_user(
            username="inactive-operator", password="long-test-password", is_active=False
        )
        cls.inactive_it_user.groups.add(cls.it_group)
        cls.superuser = User.objects.create_superuser(
            username="school-admin", email="admin@example.test", password="long-test-password"
        )

        session = AcademicSession.objects.create(
            label="2025/2026", start_date=datetime.date(2025, 9, 1),
            end_date=datetime.date(2026, 7, 31), is_current=True,
        )
        cls.term = Term.objects.create(session=session, name="first", is_current=True)
        cls.classroom = ClassRoom.objects.create(name="JSS1A", level="jss1", arm="A")
        cls.other_classroom = ClassRoom.objects.create(name="JSS1B", level="jss1", arm="B")
        cls.student = Student.objects.create(
            admission_number="TEST/JSS1/0001", first_name="Ada", last_name="Nwosu",
            current_class=cls.classroom,
        )
        cls.other_student = Student.objects.create(
            admission_number="TEST/JSS1/0002", first_name="Tunde", last_name="Adebayo",
            current_class=cls.other_classroom,
        )
        cls.draft_student = Student.objects.create(
            admission_number="TEST/JSS1/0003", first_name="Draft", last_name="Result",
            current_class=cls.classroom,
        )
        cls.blocked_student = Student.objects.create(
            admission_number="TEST/JSS1/0004", first_name="Blocked", last_name="Result",
            current_class=cls.classroom,
        )
        department = Department.objects.create(name="Dashboard Test Department")
        subject = Subject.objects.create(
            name="English", department=department, level="junior"
        )
        ResultEntry.objects.create(
            student=cls.student, term=cls.term, subject=subject, ca_score=24, exam_score=61
        )
        TermResult.objects.create(
            student=cls.student, term=cls.term, overall_total=85, average=85,
            position_in_class="1st out of 1", is_published=True,
        )
        TermResult.objects.create(
            student=cls.other_student, term=cls.term, overall_total=70, average=70,
            is_published=True,
        )
        TermResult.objects.create(
            student=cls.draft_student, term=cls.term, overall_total=60, average=60,
            is_published=False,
        )
        TermResult.objects.create(
            student=cls.blocked_student, term=cls.term, overall_total=55, average=55,
            is_published=True, is_blocked=True,
        )

    def test_staff_login_page_renders_and_login_enters_dashboard(self):
        response = self.client.get(reverse("results:staff_login"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Staff Login")

        response = self.client.post(reverse("results:staff_login"), {
            "username": "it-operator",
            "password": "long-test-password",
        })
        self.assertRedirects(response, reverse("results:staff_dashboard"))
        dashboard = self.client.get(reverse("results:staff_dashboard"))
        self.assertEqual(dashboard.status_code, 200)
        self.assertContains(dashboard, "Welcome, it-operator")

    def test_login_without_it_permission_redirects_to_access_denied(self):
        response = self.client.post(reverse("results:staff_login"), {
            "username": "ordinary-user",
            "password": "long-test-password",
        })
        self.assertRedirects(
            response, reverse("results:access_denied"), fetch_redirect_response=False
        )

    def test_logout_requires_csrf_protected_post(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.it_user)
        dashboard = client.get(reverse("results:staff_dashboard"))
        self.assertEqual(dashboard.status_code, 200)
        csrf_token = client.cookies["csrftoken"].value

        rejected = client.post(reverse("results:staff_logout"))
        self.assertEqual(rejected.status_code, 403)

        logged_out = client.post(
            reverse("results:staff_logout"),
            HTTP_X_CSRFTOKEN=csrf_token,
        )
        self.assertEqual(logged_out.status_code, 302)
        self.assertEqual(logged_out["Location"], reverse("pages:home"))

    def test_anonymous_dashboard_request_redirects_to_staff_login(self):
        response = self.client.get(reverse("results:staff_dashboard"))
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("results:staff_login"), response["Location"])

    def test_authenticated_user_without_it_permission_gets_access_denied(self):
        self.client.force_login(self.unauthorized_user)
        response = self.client.get(reverse("results:staff_dashboard"))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response["Location"], reverse("results:access_denied"))
        denied = self.client.get(reverse("results:access_denied"))
        self.assertEqual(denied.status_code, 403)

    def test_inactive_it_account_cannot_use_an_existing_session(self):
        self.client.force_login(self.inactive_it_user)
        response = self.client.get(reverse("results:staff_dashboard"))
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("results:staff_login"), response["Location"])

    def test_it_user_can_search_and_filter_students_and_results(self):
        self.client.force_login(self.it_user)
        response = self.client.get(reverse("results:staff_dashboard"), {
            "search": "TEST/JSS1",
            "classroom": self.classroom.pk,
            "session": self.term.session_id,
            "term": self.term.pk,
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Nwosu Ada")
        self.assertContains(response, "Published")
        self.assertNotContains(response, "Adebayo Tunde")
        self.assertContains(response, reverse(
            "results:staff_result_print", args=[self.student.pk, self.term.pk]
        ))

    def test_it_user_can_view_and_print_official_result_sheet(self):
        self.client.force_login(self.it_user)
        response = self.client.get(reverse(
            "results:staff_result_print", args=[self.student.pk, self.term.pk]
        ))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Mater Domini")
        self.assertContains(response, "English")
        self.assertContains(response, "24.0")
        self.assertContains(response, "61.0")

    def test_blocked_result_is_not_available_for_individual_print(self):
        from apps.results.models import TermResult

        TermResult.objects.filter(student=self.student, term=self.term).update(is_blocked=True)
        self.client.force_login(self.it_user)
        response = self.client.get(reverse(
            "results:staff_result_print", args=[self.student.pk, self.term.pk]
        ))
        self.assertEqual(response.status_code, 404)

        dashboard = self.client.get(reverse("results:staff_dashboard"), {
            "classroom": self.classroom.pk,
            "session": self.term.session_id,
            "term": self.term.pk,
        })
        self.assertContains(dashboard, "Blocked")
        self.assertNotContains(dashboard, "View / Print")

        class_print = self.client.get(reverse("results:staff_class_results_print"), {
            "classroom": self.classroom.pk,
            "session": self.term.session_id,
            "term": self.term.pk,
        })
        self.assertEqual(class_print.status_code, 200)
        self.assertNotContains(class_print, "Nwosu Ada")

    def test_draft_result_is_not_available_for_individual_print_or_dashboard_action(self):
        from apps.results.models import TermResult

        self.client.force_login(self.it_user)
        response = self.client.get(reverse(
            "results:staff_result_print", args=[self.draft_student.pk, self.term.pk]
        ))
        self.assertEqual(response.status_code, 404)

        dashboard = self.client.get(reverse("results:staff_dashboard"), {
            "classroom": self.classroom.pk,
            "session": self.term.session_id,
            "term": self.term.pk,
        })
        self.assertContains(dashboard, "Draft")
        self.assertNotContains(dashboard, reverse(
            "results:staff_result_print", args=[self.draft_student.pk, self.term.pk]
        ))

    def test_unauthorized_user_cannot_access_individual_or_class_print(self):
        self.client.force_login(self.unauthorized_user)
        individual = self.client.get(reverse(
            "results:staff_result_print", args=[self.student.pk, self.term.pk]
        ))
        class_results = self.client.get(reverse("results:staff_class_results_print"))

        self.assertEqual(individual.status_code, 302)
        self.assertEqual(individual["Location"], reverse("results:access_denied"))
        self.assertEqual(class_results.status_code, 302)
        self.assertEqual(class_results["Location"], reverse("results:access_denied"))

    def test_class_print_includes_only_students_with_term_results(self):
        self.client.force_login(self.it_user)
        response = self.client.get(reverse("results:staff_class_results_print"), {
            "classroom": self.classroom.pk,
            "session": self.term.session_id,
            "term": self.term.pk,
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Nwosu Ada")
        self.assertNotContains(response, "Adebayo Tunde")
        self.assertNotContains(response, "Draft Result")
        self.assertNotContains(response, "Blocked Result")
        self.assertContains(response, "Print All (1 students)")

    def test_it_user_cannot_access_django_admin_but_superuser_can(self):
        self.assertFalse(self.it_user.is_staff)
        self.client.force_login(self.it_user)
        dashboard = self.client.get(reverse("results:staff_dashboard"))
        self.assertContains(dashboard, "Admin Panel")
        self.assertContains(dashboard, f'href="{reverse("admin:index")}"')
        denied = self.client.get("/admin/")
        self.assertNotEqual(denied.status_code, 200)

        self.client.force_login(self.superuser)
        admin_response = self.client.get("/admin/")
        self.assertEqual(admin_response.status_code, 200)

    def test_old_scratch_card_endpoint_is_not_available(self):
        response = self.client.get("/check-result/")
        self.assertEqual(response.status_code, 404)

    def test_public_staff_login_link_uses_the_dedicated_portal(self):
        response = self.client.get(reverse("pages:home"))
        self.assertContains(response, f'href="{reverse("results:staff_login")}"')
        self.assertNotContains(response, 'href="/admin/login/"')


class StaffOnlyResultAccessTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        from django.contrib.auth.models import Group, User
        from apps.results.models import TermResult

        cls.teacher = User.objects.create_user(
            username="teacher", password="long-test-password"
        )
        cls.teacher.groups.add(Group.objects.get(name="IT Staff"))
        cls.student_user = User.objects.create_user(
            username="nonstaff", password="long-test-password", is_staff=False
        )
        session = AcademicSession.objects.create(
            label="2025/2026", start_date=datetime.date(2025, 9, 1),
            end_date=datetime.date(2026, 7, 31),
        )
        cls.term = Term.objects.create(session=session, name="first")
        classroom = ClassRoom.objects.create(name="JSS1A", level="jss1", arm="A")
        cls.student = Student.objects.create(
            admission_number="TEST/0001", first_name="Test", last_name="Student",
            current_class=classroom,
        )
        TermResult.objects.create(
            student=cls.student, term=cls.term, is_published=True
        )

    def test_old_public_result_endpoint_is_not_available(self):
        response = self.client.get("/check-result/")
        self.assertEqual(response.status_code, 404)

    def test_anonymous_user_cannot_print_a_student_result(self):
        response = self.client.get(reverse(
            "results:staff_result_print", args=[self.student.id, self.term.id]
        ))
        self.assertEqual(response.status_code, 302)

    def test_authenticated_nonstaff_user_cannot_print_a_student_result(self):
        self.client.force_login(self.student_user)
        response = self.client.get(reverse(
            "results:staff_result_print", args=[self.student.id, self.term.id]
        ))
        self.assertNotEqual(response.status_code, 200)

    def test_staff_user_can_print_a_student_result(self):
        self.client.force_login(self.teacher)
        response = self.client.get(reverse(
            "results:staff_result_print", args=[self.student.id, self.term.id]
        ))
        self.assertEqual(response.status_code, 200)


class StaffToolsTests(TestCase):
    """Covers CSV bulk upload, subject-level disambiguation, the master
    sheet, and the shared position/average calculation service."""

    @classmethod
    def setUpTestData(cls):
        from django.contrib.auth.models import User
        from apps.results.models import ResultEntry, TermResult

        cls.staff_user = User.objects.create_user(
            username="staffmember", password="testpass123", is_staff=True
        )
        from django.contrib.auth.models import Group
        cls.staff_user.groups.add(Group.objects.get(name="IT Staff"))

        session = AcademicSession.objects.create(
            label="2025/2026", start_date=datetime.date(2025, 9, 1),
            end_date=datetime.date(2026, 7, 31), is_current=True,
        )
        cls.term = Term.objects.create(session=session, name="first", is_current=True)
        cls.classroom = ClassRoom.objects.create(name="JSS1A", level="jss1", arm="A")

        department = Department.objects.create(name="Core")
        # Deliberately create "Mathematics" at BOTH levels — the exact
        # ambiguity that caused a real bug during development.
        cls.junior_maths = Subject.objects.create(name="Mathematics", department=department, level="junior")
        cls.senior_maths = Subject.objects.create(name="Mathematics", department=department, level="senior")

        cls.student = Student.objects.create(
            admission_number="TEST/JSS1/0001", first_name="Jane", last_name="Doe",
            current_class=cls.classroom,
        )

    def _csv_upload(self, csv_bytes, filename="upload.csv"):
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.login(username="staffmember", password="testpass123")
        upload = SimpleUploadedFile(filename, csv_bytes, content_type="text/csv")
        return self.client.post(reverse("results:upload_scores"), {
            "term": self.term.id, "csv_file": upload,
        })

    def test_anonymous_user_cannot_access_upload_page(self):
        response = self.client.get(reverse("results:upload_scores"))
        self.assertNotEqual(response.status_code, 200)

    def test_staff_can_access_upload_page(self):
        self.client.login(username="staffmember", password="testpass123")
        response = self.client.get(reverse("results:upload_scores"))
        self.assertEqual(response.status_code, 200)

    def test_csv_upload_resolves_ambiguous_subject_by_class_level(self):
        """The core bug this test suite exists to prevent regressing: a JSS1
        student's 'Mathematics' score must land on the JUNIOR Subject record,
        never the senior one, even though both share the exact same name."""
        csv_bytes = b"admission_number,subject,ca_score,exam_score\nTEST/JSS1/0001,Mathematics,25,60\n"
        response = self._csv_upload(csv_bytes)
        self.assertContains(response, "Upload complete")

        from apps.results.models import ResultEntry
        entry = ResultEntry.objects.get(student=self.student, term=self.term)
        self.assertEqual(entry.subject_id, self.junior_maths.id)
        self.assertNotEqual(entry.subject_id, self.senior_maths.id)

    def test_csv_upload_reports_unknown_admission_number_without_crashing(self):
        csv_bytes = b"admission_number,subject,ca_score,exam_score\nNOTREAL/0000,Mathematics,25,60\n"
        response = self._csv_upload(csv_bytes)
        self.assertContains(response, "no student with admission number")

    def test_csv_upload_rejects_out_of_range_scores(self):
        csv_bytes = b"admission_number,subject,ca_score,exam_score\nTEST/JSS1/0001,Mathematics,99,60\n"
        response = self._csv_upload(csv_bytes)
        self.assertContains(response, "out of range")

    def test_csv_upload_recalculates_totals_and_positions(self):
        csv_bytes = b"admission_number,subject,ca_score,exam_score\nTEST/JSS1/0001,Mathematics,25,60\n"
        self._csv_upload(csv_bytes)

        from apps.results.models import TermResult
        term_result = TermResult.objects.get(student=self.student, term=self.term)
        self.assertEqual(term_result.overall_total, 85)
        self.assertEqual(term_result.position_in_class, "1st out of 1")

    def test_master_sheet_requires_staff_login(self):
        response = self.client.get(reverse("results:master_sheet"))
        self.assertNotEqual(response.status_code, 200)

    def test_master_sheet_renders_with_class_and_term_selected(self):
        self._csv_upload(b"admission_number,subject,ca_score,exam_score\nTEST/JSS1/0001,Mathematics,25,60\n")
        self.client.login(username="staffmember", password="testpass123")
        response = self.client.get(reverse("results:master_sheet"), {
            "term": self.term.id, "classroom": self.classroom.id,
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Doe Jane")
        self.assertContains(response, "85.0")  # the total we uploaded


class RecalculatePositionsServiceTests(TestCase):
    """Directly tests the shared ranking service in isolation from any view."""

    def test_positions_never_cross_class_boundaries(self):
        from apps.results import services
        from apps.results.models import ResultEntry, TermResult

        session = AcademicSession.objects.create(
            label="2025/2026", start_date=datetime.date(2025, 9, 1),
            end_date=datetime.date(2026, 7, 31), is_current=True,
        )
        term = Term.objects.create(session=session, name="first", is_current=True)
        department = Department.objects.create(name="Core")
        subject = Subject.objects.create(name="English", department=department, level="junior")

        class_a = ClassRoom.objects.create(name="JSS1A", level="jss1", arm="A")
        class_b = ClassRoom.objects.create(name="JSS1B", level="jss1", arm="B")

        # Weakest student in Class A scores higher than the strongest in Class B.
        weak_in_a = Student.objects.create(admission_number="A/1", first_name="Weak", last_name="InA", current_class=class_a)
        strong_in_a = Student.objects.create(admission_number="A/2", first_name="Strong", last_name="InA", current_class=class_a)
        strong_in_b = Student.objects.create(admission_number="B/1", first_name="Strong", last_name="InB", current_class=class_b)

        ResultEntry.objects.create(student=weak_in_a, term=term, subject=subject, ca_score=20, exam_score=40)
        ResultEntry.objects.create(student=strong_in_a, term=term, subject=subject, ca_score=30, exam_score=70)
        ResultEntry.objects.create(student=strong_in_b, term=term, subject=subject, ca_score=10, exam_score=20)

        for student in [weak_in_a, strong_in_a, strong_in_b]:
            services.recalculate_term_result_totals(student, term)
        services.recalculate_positions_for_term(term)

        weak_entry = ResultEntry.objects.get(student=weak_in_a, term=term, subject=subject)
        strong_entry = ResultEntry.objects.get(student=strong_in_a, term=term, subject=subject)

        # Within Class A, weak_in_a must be LAST (2nd out of 2 in that class),
        # never boosted by comparison against the lower-scoring Class B student.
        self.assertEqual(weak_entry.position_in_subject, "2nd")
        self.assertEqual(strong_entry.position_in_subject, "1st")

        # Class B's single student must be ranked only within Class B (1st
        # out of 1), not penalized for scoring lower than Class A students.
        b_term_result = TermResult.objects.get(student=strong_in_b, term=term)
        self.assertEqual(b_term_result.position_in_class, "1st out of 1")


class WorkbookUploadTests(TestCase):
    """
    Covers the Excel workbook importer (REGISTER + term sheets), including
    the class-level bug we found and fixed: ClassRoom.level must be stored
    as 'jss1'/'ss2' (matching LEVEL_CHOICES), never 'junior'/'senior' —
    getting this wrong silently breaks subject-level disambiguation.
    """

    @classmethod
    def setUpTestData(cls):
        from django.contrib.auth.models import User
        cls.staff_user = User.objects.create_user(
            username="staffuser", password="testpass123", is_staff=True
        )
        cls.session = AcademicSession.objects.create(
            label="2025/2026", start_date=datetime.date(2025, 9, 1),
            end_date=datetime.date(2026, 7, 31), is_current=True,
        )
        department = Department.objects.create(name="Test Dept")
        Subject.objects.create(name="English Studies", department=department, level="junior")
        Subject.objects.create(name="Mathematics", department=department, level="junior")
        Subject.objects.create(name="Mathematics", department=department, level="senior")

    def _build_workbook(self):
        import io
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "REGISTER"
        ws["A2"] = "Class Code:"
        ws["B2"] = "J1A"
        ws["A4"] = "S/N"; ws["B4"] = "Admission No."; ws["C4"] = "Surname"; ws["D4"] = "First Name"
        ws["A5"] = 1; ws["B5"] = "MDSJ1A001"; ws["C5"] = "Nwosu"; ws["D5"] = "Adaeze"

        ws2 = wb.create_sheet("1ST TERM")
        ws2["A4"] = "Admission No."; ws2["B4"] = "Name"
        ws2["C3"] = "English Studies"; ws2["G3"] = "Mathematics"
        ws2["C4"] = "CA"; ws2["D4"] = "Exam"; ws2["E4"] = "Total"; ws2["F4"] = "Grade"
        ws2["G4"] = "CA"; ws2["H4"] = "Exam"; ws2["I4"] = "Total"; ws2["J4"] = "Grade"
        ws2["A5"] = "MDSJ1A001"; ws2["B5"] = "Nwosu Adaeze"
        ws2["C5"] = 24; ws2["D5"] = 60
        ws2["G5"] = 20; ws2["H5"] = 58

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.read()

    def test_anonymous_access_blocked(self):
        response = self.client.get(reverse("results:upload_workbook"))
        self.assertNotEqual(response.status_code, 200)

    def test_workbook_upload_creates_student_and_classroom_correctly(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        self.client.force_login(self.staff_user)
        upload = SimpleUploadedFile(
            "test.xlsx", self._build_workbook(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response = self.client.post(reverse("results:upload_workbook"), {
            "session": self.session.id, "workbook_file": upload,
        })
        self.assertEqual(response.status_code, 200)

        student = Student.objects.get(admission_number="MDSJ1A001")
        self.assertEqual(student.first_name, "Adaeze")
        self.assertEqual(student.last_name, "Nwosu")

        # The bug we caught: level must be 'jss1', not 'junior'.
        classroom = ClassRoom.objects.get(name="JSS1A")
        self.assertEqual(classroom.level, "jss1")
        self.assertEqual(classroom.subject_level, "junior")

    def test_workbook_upload_writes_correct_scores_to_correct_subject_level(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from apps.results.models import ResultEntry, Term

        self.client.force_login(self.staff_user)
        upload = SimpleUploadedFile(
            "test.xlsx", self._build_workbook(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.client.post(reverse("results:upload_workbook"), {
            "session": self.session.id, "workbook_file": upload,
        })

        term = Term.objects.get(session=self.session, name="first")
        student = Student.objects.get(admission_number="MDSJ1A001")
        math_entry = ResultEntry.objects.get(student=student, term=term, subject__name="Mathematics")

        # Must resolve to the JUNIOR Mathematics, not the senior one —
        # this is the ambiguity the level-scoped lookup exists to prevent.
        self.assertEqual(math_entry.subject.level, "junior")
        self.assertEqual(math_entry.ca_score, 20)
        self.assertEqual(math_entry.exam_score, 58)


@override_settings(STORAGES={
    "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
    "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
})
class StaffPrintViewTests(TestCase):
    """
    Covers the print-without-scratch-card staff views added for form
    teachers/admin — must require staff login, and must actually show
    the requested student/class's real data.
    """

    @classmethod
    def setUpTestData(cls):
        from django.contrib.auth.models import Group, User
        from apps.results.models import ResultEntry, TermResult

        cls.staff_user = User.objects.create_user(
            username="printstaff", password="testpass123"
        )
        cls.staff_user.groups.add(Group.objects.get(name="IT Staff"))
        session = AcademicSession.objects.create(
            label="2025/2026", start_date=datetime.date(2025, 9, 1),
            end_date=datetime.date(2026, 7, 31), is_current=True,
        )
        cls.term = Term.objects.create(session=session, name="first", is_current=True)
        cls.classroom = ClassRoom.objects.create(name="JSS1A", level="jss1", arm="A")
        cls.student = Student.objects.create(
            admission_number="PRINT/0001", first_name="Print", last_name="Test",
            current_class=cls.classroom,
        )
        department = Department.objects.create(name="Print Test Dept")
        subject = Subject.objects.create(name="Print Subject", department=department, level="junior")
        ResultEntry.objects.create(student=cls.student, term=cls.term, subject=subject, ca_score=20, exam_score=50)
        cls.term_result = TermResult.objects.create(
            student=cls.student, term=cls.term, overall_total=70, average=70, is_published=True,
        )

    def test_individual_print_requires_staff_login(self):
        url = reverse("results:staff_result_print", args=[self.student.id, self.term.id])
        response = self.client.get(url, follow=True)
        self.assertIn("login", response.request["PATH_INFO"].lower())

    def test_individual_print_shows_correct_student_when_logged_in(self):
        self.client.force_login(self.staff_user)
        url = reverse("results:staff_result_print", args=[self.student.id, self.term.id])
        response = self.client.get(url)
        self.assertContains(response, "Test Print")  # full_name is "Last First"
        self.assertContains(response, "Print Subject")

    def test_individual_print_does_not_require_scratch_card(self):
        """The whole point of this view: no PIN/serial needed at all."""
        self.client.force_login(self.staff_user)
        url = reverse("results:staff_result_print", args=[self.student.id, self.term.id])
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)

    def test_class_bulk_print_requires_staff_login(self):
        response = self.client.get(reverse("results:staff_class_results_print"), follow=True)
        self.assertIn("login", response.request["PATH_INFO"].lower())

    def test_class_bulk_print_lists_every_student_in_class(self):
        self.client.force_login(self.staff_user)
        url = (
            f"{reverse('results:staff_class_results_print')}?session={self.term.session_id}"
            f"&term={self.term.id}&classroom={self.classroom.id}"
        )
        response = self.client.get(url)
        self.assertContains(response, "Test Print")
