import csv
import io
from functools import wraps

import openpyxl
from django.contrib.auth.views import LoginView, LogoutView, redirect_to_login
from django.contrib.admin.views.decorators import staff_member_required
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Prefetch, Q
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils.decorators import method_decorator
from django.views import View
from django.views.generic import TemplateView

from . import services
from .forms import (
    MasterSheetForm,
    PromotionResultForm,
    ScoreCSVUploadForm,
    StaffDashboardFilterForm,
    WorkbookUploadForm,
)
from .models import (
    AcademicSession,
    ClassRoom,
    ResultEntry,
    PsychomotorRating,
    SessionResult,
    Student,
    Term,
    TermResult,
)
from apps.academics.models import Subject

IT_STAFF_PERMISSION = "results.access_student_results"


def it_staff_required(view_func):
    @wraps(view_func)
    def wrapped(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return redirect_to_login(
                request.get_full_path(), login_url=reverse("results:staff_login")
            )
        if request.user.is_active and (
            request.user.is_superuser or request.user.has_perm(IT_STAFF_PERMISSION)
        ):
            return view_func(request, *args, **kwargs)
        return redirect("results:access_denied")

    return wrapped


class StaffLoginView(LoginView):
    template_name = "results/staff/login.html"
    redirect_authenticated_user = True

    def get_success_url(self):
        if self.request.user.is_active and (
            self.request.user.is_superuser or self.request.user.has_perm(IT_STAFF_PERMISSION)
        ):
            return self.get_redirect_url() or reverse("results:staff_dashboard")
        return reverse("results:access_denied")


class StaffLogoutView(LogoutView):
    next_page = "pages:home"


class StaffAccessDeniedView(TemplateView):
    template_name = "results/staff/access_denied.html"

    def get(self, request, *args, **kwargs):
        return self.render_to_response(self.get_context_data(), status=403)


@method_decorator(it_staff_required, name="dispatch")
class StaffDashboardView(View):
    template_name = "results/staff/dashboard.html"

    def get(self, request):
        current_term = Term.objects.filter(is_current=True).select_related("session").first()
        initial = {}
        if current_term:
            initial = {"session": current_term.session_id, "term": current_term.pk}

        form = StaffDashboardFilterForm(request.GET or None, initial=initial)
        if request.GET:
            filters_valid = form.is_valid()
            filters = form.cleaned_data if filters_valid else {}
        else:
            filters_valid = True
            filters = initial

        students = Student.objects.filter(is_active=True).select_related("current_class")
        search = filters.get("search", "")
        if search:
            students = students.filter(
                Q(first_name__icontains=search)
                | Q(middle_name__icontains=search)
                | Q(last_name__icontains=search)
                | Q(admission_number__icontains=search)
            )
        if filters.get("classroom"):
            students = students.filter(current_class_id=filters["classroom"])

        selected_term_id = filters.get("term")
        selected_term = Term.objects.filter(pk=selected_term_id).first() if selected_term_id else None
        result_queryset = TermResult.objects.none()
        if selected_term:
            result_queryset = TermResult.objects.filter(term=selected_term).select_related("term")

        students = students.prefetch_related(
            Prefetch("term_results", queryset=result_queryset, to_attr="dashboard_term_results")
        ).order_by("last_name", "first_name")
        page = Paginator(students, 30).get_page(request.GET.get("page"))
        for student in page.object_list:
            student.dashboard_result = next(iter(student.dashboard_term_results), None)

        return render(request, self.template_name, {
            "form": form,
            "filters_valid": filters_valid,
            "students_page": page,
            "selected_term": selected_term,
            "selected_term_id": selected_term.pk if selected_term else "",
            "selected_session_id": filters.get("session", ""),
            "selected_classroom_id": filters.get("classroom", ""),
            "active_student_count": Student.objects.filter(is_active=True).count(),
            "class_count": ClassRoom.objects.count(),
            "session_count": AcademicSession.objects.count(),
            "result_count": TermResult.objects.count(),
            "staff_name": request.user.get_full_name() or request.user.get_username(),
        })


# ==============================================================================
# Staff-only tools: CSV score upload, position recalculation, master sheets
# ==============================================================================
@method_decorator(staff_member_required, name="dispatch")
class ScoreUploadView(View):
    """
    Bulk score upload for a whole class/term at once from a CSV file,
    rather than entering each ResultEntry one at a time in the admin.

    Gated with @staff_member_required (Django's built-in admin-login check)
    rather than building separate authentication — this reuses the same
    login every other admin action already requires.
    """
    template_name = "results/staff/upload_scores.html"

    def get(self, request):
        return render(request, self.template_name, {"form": ScoreCSVUploadForm()})

    def post(self, request):
        form = ScoreCSVUploadForm(request.POST, request.FILES)
        if not form.is_valid():
            return render(request, self.template_name, {"form": form})

        term = Term.objects.get(pk=form.cleaned_data["term"])
        csv_file = form.cleaned_data["csv_file"]

        decoded = csv_file.read().decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(decoded))

        created, updated, errors = 0, 0, []
        touched_students = set()

        required_columns = {"admission_number", "subject", "ca_score", "exam_score"}
        if not required_columns.issubset(set(reader.fieldnames or [])):
            errors.append(
                f"CSV header must include: {', '.join(sorted(required_columns))}. "
                f"Found: {', '.join(reader.fieldnames or []) or '(empty file)'}"
            )
            return render(request, self.template_name, {"form": form, "errors": errors})

        for row_number, row in enumerate(reader, start=2):  # row 1 is the header
            admission_number = (row.get("admission_number") or "").strip()
            subject_name = (row.get("subject") or "").strip()

            student = Student.objects.filter(admission_number__iexact=admission_number).first()
            if not student:
                errors.append(f"Row {row_number}: no student with admission number '{admission_number}'.")
                continue

            # Subject names like "Mathematics" exist at BOTH junior and
            # senior level — resolve using the student's own class level so
            # a JSS1 student's "Mathematics" never gets attached to the
            # senior-level Subject record (or vice versa).
            subject_qs = Subject.objects.filter(name__iexact=subject_name)
            if student.current_class:
                subject = subject_qs.filter(level=student.current_class.subject_level).first()
                if not subject:
                    # Fall back to any match (e.g. the subject only exists
                    # at one level) rather than failing outright.
                    subject = subject_qs.first()
            else:
                subject = subject_qs.first()

            if not subject:
                errors.append(f"Row {row_number}: no subject named '{subject_name}'.")
                continue

            try:
                ca_score = float(row.get("ca_score") or 0)
                exam_score = float(row.get("exam_score") or 0)
            except ValueError:
                errors.append(f"Row {row_number}: CA/Exam score must be a number.")
                continue

            if ca_score > 30 or exam_score > 70 or ca_score < 0 or exam_score < 0:
                errors.append(
                    f"Row {row_number}: scores out of range (CA max 30, Exam max 70) for {admission_number}."
                )
                continue

            entry, was_created = ResultEntry.objects.update_or_create(
                student=student, term=term, subject=subject,
                defaults={"ca_score": ca_score, "exam_score": exam_score},
            )
            created += 1 if was_created else 0
            updated += 0 if was_created else 1
            touched_students.add(student.id)

        for student_id in touched_students:
            services.recalculate_term_result_totals(Student.objects.get(pk=student_id), term)
        entries_updated, term_results_updated = services.recalculate_positions_for_term(term)

        return render(request, self.template_name, {
            "form": ScoreCSVUploadForm(),
            "summary": {
                "created": created,
                "updated": updated,
                "students_touched": len(touched_students),
                "entries_ranked": entries_updated,
                "term_results_ranked": term_results_updated,
            },
            "errors": errors,
        })


@method_decorator(it_staff_required, name="dispatch")
class MasterSheetView(View):
    """
    Printable class-wide broadsheet: every active student in a class, every
    subject as a column, for one term OR cumulative results for one session.
    For teachers/admin, not parents. Supports both term-based and session 
    cumulative (annual broadsheet) displays with flexible print formatting.
    """
    template_name = "results/staff/master_sheet.html"

    def get(self, request):
        from .models import AcademicSession, SessionResult
        
        result_type = request.GET.get("result_type", "term")
        classroom_id = request.GET.get("classroom")
        
        form_kwargs = {"initial": {"result_type": result_type}}
        if classroom_id:
            form_kwargs["initial"]["classroom"] = classroom_id

        if result_type == "session_cumulative":
            session_id = request.GET.get("session")
            if session_id:
                form_kwargs["initial"]["session"] = session_id
            
            if not session_id or not classroom_id:
                return render(request, self.template_name, {
                    "form": MasterSheetForm(**form_kwargs),
                    "result_type": "session_cumulative"
                })

            session = AcademicSession.objects.get(pk=session_id)
            classroom = ClassRoom.objects.get(pk=classroom_id)
            students = list(classroom.students.filter(is_active=True).order_by("last_name", "first_name"))
            
            # Get session results (cumulative annual data)
            session_results_by_student = {}
            for student in students:
                session_result = SessionResult.objects.filter(
                    student=student, session=session
                ).first()
                session_results_by_student[student.id] = session_result

            return render(request, self.template_name, {
                "form": MasterSheetForm(**form_kwargs),
                "result_type": "session_cumulative",
                "session": session,
                "classroom": classroom,
                "students": students,
                "session_results_by_student": session_results_by_student,
            })
        else:
            # Original term-based view
            term_id = request.GET.get("term")
            if term_id:
                form_kwargs["initial"]["term"] = term_id
            
            if not term_id or not classroom_id:
                return render(request, self.template_name, {
                    "form": MasterSheetForm(**form_kwargs),
                    "result_type": "term"
                })

            term = Term.objects.get(pk=term_id)
            classroom = ClassRoom.objects.get(pk=classroom_id)
            students = list(classroom.students.filter(is_active=True).order_by("last_name", "first_name"))

            subjects = list(
                Subject.objects.filter(
                    result_entries__student__in=students, result_entries__term=term
                ).distinct().order_by("name")
            )

            entries_by_student = {}
            for student in students:
                entries = {
                    e.subject_id: e
                    for e in ResultEntry.objects.filter(student=student, term=term)
                }
                term_result = TermResult.objects.filter(student=student, term=term).first()
                entries_by_student[student.id] = {"entries": entries, "term_result": term_result}

            return render(request, self.template_name, {
                "form": MasterSheetForm(**form_kwargs),
                "result_type": "term",
                "term": term,
                "classroom": classroom,
                "students": students,
                "subjects": subjects,
                "entries_by_student": entries_by_student,
            })


TERM_SHEET_NAMES = {"1ST TERM": "first", "2ND TERM": "second", "3RD TERM": "third"}
SUBJECT_BLOCK_START_COL = 2  # 0-indexed: col 0=admission_no, col 1=student name, col 2 onward=subjects
SUBJECT_BLOCK_WIDTH = 4      # CA, Exam, Total, Grade


@method_decorator(staff_member_required, name="dispatch")
class WorkbookUploadView(View):
    """
    Uploads the school's real result workbook (REGISTER + 1ST/2ND/3RD TERM +
    BROADSHEET sheets) for one class/arm at a time. In one pass:
      1. REGISTER  -> creates any Student who doesn't already exist
      2. each TERM sheet present -> creates/updates ResultEntry rows,
         skipping subjects with no score entered (a school uploading
         mid-term won't have every subject filled in yet)
      3. BROADSHEET -> creates/updates the SessionResult (annual summary)

    Only step 1 requires the REGISTER sheet; any of the other sheets can
    be absent (e.g. uploading First Term only, before Second/Third exist).
    """
    template_name = "results/staff/upload_workbook.html"

    def get(self, request):
        return render(request, self.template_name, {"form": WorkbookUploadForm()})

    def post(self, request):
        form = WorkbookUploadForm(request.POST, request.FILES)
        if not form.is_valid():
            return render(request, self.template_name, {"form": form})

        session = AcademicSession.objects.get(pk=form.cleaned_data["session"])
        workbook_file = form.cleaned_data["workbook_file"]

        try:
            wb = openpyxl.load_workbook(workbook_file, data_only=True)
        except Exception as exc:
            return render(request, self.template_name, {
                "form": form, "errors": [f"Could not read this file as an Excel workbook: {exc}"],
            })

        errors = []
        summary = {"students_created": 0, "students_updated": 0, "scores_written": 0, "terms_processed": []}

        classroom, class_errors = self._process_register(wb, summary)
        errors.extend(class_errors)

        if classroom:
            for sheet_name, term_key in TERM_SHEET_NAMES.items():
                if sheet_name in wb.sheetnames:
                    term_errors = self._process_term_sheet(wb[sheet_name], term_key, session, classroom, summary)
                    errors.extend(term_errors)
                    if not term_errors:
                        summary["terms_processed"].append(sheet_name)

            if "BROADSHEET" in wb.sheetnames:
                broadsheet_errors = self._process_broadsheet(wb["BROADSHEET"], session, summary)
                errors.extend(broadsheet_errors)

        return render(request, self.template_name, {
            "form": WorkbookUploadForm(),
            "summary": summary if classroom else None,
            "classroom": classroom,
            "errors": errors,
        })

    # ------------------------------------------------------------------
    def _process_register(self, wb, summary):
        """Parses the REGISTER sheet: identifies the class/arm, creates any
        missing Student records. Returns (classroom_or_None, errors)."""
        if "REGISTER" not in wb.sheetnames:
            return None, ["This workbook has no REGISTER sheet — cannot identify the class or students."]

        ws = wb["REGISTER"]
        errors = []

        # Class code is a labeled cell, e.g. "Class Code:" | "J1A" — search
        # the first few rows for it rather than assuming a fixed cell,
        # since the register's header layout could shift slightly.
        class_code = None
        for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
            for i, cell in enumerate(row):
                if cell and str(cell).strip().lower().startswith("class code"):
                    if i + 1 < len(row) and row[i + 1]:
                        class_code = str(row[i + 1]).strip()
        if not class_code:
            return None, ["Could not find 'Class Code:' in the REGISTER sheet."]

        classroom_name, level = services.classroom_code_to_name(class_code)
        if not classroom_name:
            return None, [f"Class code '{class_code}' doesn't match the expected pattern (e.g. J1A, S2B)."]

        classroom, _ = ClassRoom.objects.get_or_create(
            name=classroom_name,
            defaults={"level": classroom_name[:-1].lower(), "arm": classroom_name[-1]},
        )

        # Find the header row (contains "Admission No." literally), then read
        # every row below it until a blank admission number ends the list.
        header_row_index = None
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
            if any(cell and "admission" in str(cell).lower() for cell in row):
                header_row_index = row_idx
                break
        if not header_row_index:
            return classroom, ["Could not find the 'Admission No.' header row in REGISTER."]

        for row in ws.iter_rows(min_row=header_row_index + 1, values_only=True):
            admission_number = (str(row[1]).strip() if row[1] else "")
            if not admission_number:
                continue
            surname = (str(row[2]).strip() if row[2] else "")
            first_name = (str(row[3]).strip() if row[3] else "")
            if not surname or not first_name:
                errors.append(f"Register row for '{admission_number}' is missing a surname or first name — skipped.")
                continue

            student, was_created = Student.objects.get_or_create(
                admission_number=admission_number,
                defaults={"first_name": first_name, "last_name": surname, "current_class": classroom},
            )
            if was_created:
                summary["students_created"] += 1
            else:
                # Keep an existing student's class in sync with this register.
                if student.current_class_id != classroom.id:
                    student.current_class = classroom
                    student.save(update_fields=["current_class"])
                summary["students_updated"] += 1

        return classroom, errors

    @staticmethod
    def _level_code(classroom_name):
        return "junior" if classroom_name.startswith("JSS") else "senior"

    # ------------------------------------------------------------------
    def _process_term_sheet(self, ws, term_key, session, classroom, summary):
        """Import one term using cached lookups and batched score writes."""
        from decimal import Decimal, InvalidOperation

        errors = []
        term, _ = Term.objects.get_or_create(session=session, name=term_key)
        level = self._level_code(classroom.name)

        # Resolve each subject header once, rather than querying for every row.
        header_row = [cell.value for cell in ws[3]]
        subject_columns = []
        col = SUBJECT_BLOCK_START_COL
        stop_headers = {"OVERALL TOTAL", "SUBJECT EVALUATED", "SUBJECT OFFERED", "AVERAGE", "POSITION"}
        while col < len(header_row) and header_row[col]:
            header_text = str(header_row[col]).strip()
            if header_text.upper() in stop_headers:
                break
            subject = services.resolve_subject(header_text, level)
            if subject is None:
                errors.append(f"{term.get_name_display()}: unrecognized subject column '{header_text}' — skipped.")
            else:
                subject_columns.append((subject, col))
            col += SUBJECT_BLOCK_WIDTH

        # Load students once; admission numbers are matched case-insensitively.
        students_by_admission = {
            s.admission_number.casefold(): s
            for s in Student.objects.all()
        }
        # Keep support for previously registered students whose current class
        # may not have been synchronized yet, without one query per workbook row.
        admission_values = set()
        parsed_rows = []
        for row_number, row in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
            admission = str(row[0]).strip() if row and row[0] else ""
            if not admission:
                continue
            admission_values.add(admission.casefold())
            parsed_rows.append((row_number, admission, row))

        # Validate and collect changes before touching result rows.
        pending = {}
        touched_student_ids = set()
        for row_number, admission, row in parsed_rows:
            student = students_by_admission.get(admission.casefold())
            if not student:
                errors.append(f"{term.get_name_display()}: no student '{admission}' on row {row_number} — run REGISTER import first.")
                continue
            for subject, ca_col in subject_columns:
                ca_raw = row[ca_col] if ca_col < len(row) else None
                exam_raw = row[ca_col + 1] if ca_col + 1 < len(row) else None
                if ca_raw is None and exam_raw is None:
                    continue
                try:
                    ca = Decimal(str(ca_raw or 0))
                    exam = Decimal(str(exam_raw or 0))
                except (InvalidOperation, TypeError, ValueError):
                    errors.append(f"{term.get_name_display()}: invalid score for {admission}, {subject.name} — skipped.")
                    continue
                if ca < 0 or ca > 30 or exam < 0 or exam > 70:
                    errors.append(f"{term.get_name_display()}: {admission}, {subject.name}: CA must be 0–30 and exam 0–70 — skipped.")
                    continue
                pending[(student.id, subject.id)] = (student, subject, ca, exam)
                touched_student_ids.add(student.id)

        if not pending:
            return errors

        student_ids = list(touched_student_ids)
        subject_ids = list({key[1] for key in pending})
        with transaction.atomic():
            existing = {
                (entry.student_id, entry.subject_id): entry
                for entry in ResultEntry.objects.filter(
                    term=term, student_id__in=student_ids, subject_id__in=subject_ids
                )
            }
            to_create, to_update = [], []
            for key, (student, subject, ca, exam) in pending.items():
                entry = existing.get(key)
                if entry:
                    entry.ca_score = ca
                    entry.exam_score = exam
                    to_update.append(entry)
                else:
                    to_create.append(ResultEntry(
                        student=student, term=term, subject=subject,
                        ca_score=ca, exam_score=exam,
                    ))
            if to_create:
                ResultEntry.objects.bulk_create(to_create, batch_size=500)
            if to_update:
                ResultEntry.objects.bulk_update(to_update, ["ca_score", "exam_score"], batch_size=500)
            summary["scores_written"] += len(pending)

            # Recompute totals in memory from one query, preserving publication
            # and blocking flags on existing TermResult rows.
            from collections import defaultdict
            totals = defaultdict(lambda: [Decimal("0"), 0])
            for entry in ResultEntry.objects.filter(term=term, student_id__in=student_ids):
                totals[entry.student_id][0] += (entry.ca_score or 0) + (entry.exam_score or 0)
                totals[entry.student_id][1] += 1
            current_results = {
                tr.student_id: tr for tr in TermResult.objects.filter(term=term, student_id__in=student_ids)
            }
            new_results, changed_results = [], []
            for student_id in student_ids:
                total, count = totals[student_id]
                average = round(total / count, 1) if count else None
                result = current_results.get(student_id)
                if result:
                    result.overall_total = total
                    result.average = average
                    changed_results.append(result)
                else:
                    new_results.append(TermResult(
                        student_id=student_id, term=term, overall_total=total, average=average
                    ))
            if new_results:
                TermResult.objects.bulk_create(new_results, batch_size=500)
            if changed_results:
                TermResult.objects.bulk_update(changed_results, ["overall_total", "average"], batch_size=500)

        services.recalculate_positions_for_term(term)
        return errors

    # ------------------------------------------------------------------
    def _process_broadsheet(self, ws, session, summary):
        """Parses the BROADSHEET sheet: one row per student with the
        finalized annual totals and the staff's promotion decision."""
        errors = []
        header_row_index = None
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=8, values_only=True), start=1):
            if any(cell and "admission" in str(cell).lower() for cell in row):
                header_row_index = row_idx
                break
        if not header_row_index:
            errors.append("Could not find the header row in BROADSHEET — session summary not imported.")
            return errors

        promotion_map = {
            "PROMOTED": "promoted", "PROMOTED_ON_TRIAL": "promoted_on_trial", "REPEAT": "repeat",
        }

        for row in ws.iter_rows(min_row=header_row_index + 1, values_only=True):
            admission_number = (str(row[1]).strip() if row[1] else "")
            if not admission_number:
                continue
            student = Student.objects.filter(admission_number__iexact=admission_number).first()
            if not student:
                errors.append(f"BROADSHEET: no student '{admission_number}'.")
                continue

            raw_status = (str(row[9]).strip().upper() if len(row) > 9 and row[9] else "")
            SessionResult.objects.update_or_create(
                student=student, session=session,
                defaults={
                    "first_term_total": row[3], "second_term_total": row[4], "third_term_total": row[5],
                    "cumulative_total": row[6], "session_average": row[7],
                    "overall_position": str(row[8]) if row[8] is not None else "",
                    "promotion_status": promotion_map.get(raw_status, ""),
                },
            )
        return errors


# ==============================================================================
# Staff-only printing: individual results and whole-class bulk printing,
# without a scratch card — for form teachers/admin who already have
# legitimate access to every student's record, not just their own child's.
# ==============================================================================
@method_decorator(it_staff_required, name="dispatch")
class StaffResultPrintView(View):
    """Print one student's official result sheet for authorized IT staff."""
    template_name = "results/staff/result_print.html"

    def get(self, request, student_id, term_id):
        student = get_object_or_404(Student, pk=student_id)
        term = get_object_or_404(Term, pk=term_id)
        term_result = get_object_or_404(TermResult, student=student, term=term, is_blocked=False)
        result_entries = ResultEntry.objects.filter(student=student, term=term).select_related("subject")
        psychomotor_ratings = PsychomotorRating.objects.filter(
            student=student, term=term
        ).select_related("skill")

        return render(request, self.template_name, {
            "term_result": term_result,
            "student": student,
            "result_entries": result_entries,
            "psychomotor_ratings": psychomotor_ratings,
        })


@method_decorator(it_staff_required, name="dispatch")
class StaffClassResultsPrintView(View):
    """
    Every active student's result sheet in a class, for one term, stacked
    onto one printable page with a page-break between each — this is what
    a form teacher prints once to have a paper copy of their whole class.
    """
    template_name = "results/staff/class_results_print.html"

    def get(self, request):
        form = StaffDashboardFilterForm(request.GET or None)
        if not request.GET or not form.is_valid():
            return render(request, self.template_name, {"form": form, "sheets": []})

        session_id = form.cleaned_data.get("session")
        term_id = form.cleaned_data.get("term")
        classroom_id = form.cleaned_data.get("classroom")
        if not session_id or not term_id or not classroom_id:
            form.add_error(None, "Select a class, academic session, and term to print class results.")
            return render(request, self.template_name, {"form": form, "sheets": []})

        term = get_object_or_404(Term, pk=term_id, session_id=session_id)
        classroom = get_object_or_404(ClassRoom, pk=classroom_id)
        students = classroom.students.filter(
            is_active=True,
            term_results__term=term,
            term_results__is_blocked=False,
        ).distinct().order_by("last_name", "first_name").prefetch_related(
            Prefetch(
                "term_results",
                queryset=TermResult.objects.filter(term=term, is_blocked=False),
                to_attr="selected_term_results",
            ),
            Prefetch(
                "result_entries",
                queryset=ResultEntry.objects.filter(term=term).select_related("subject"),
                to_attr="selected_result_entries",
            ),
            Prefetch(
                "psychomotor_ratings",
                queryset=PsychomotorRating.objects.filter(term=term).select_related("skill"),
                to_attr="selected_psychomotor_ratings",
            ),
        )

        sheets = [{
            "student": student,
            "term_result": student.selected_term_results[0],
            "result_entries": student.selected_result_entries,
            "psychomotor_ratings": student.selected_psychomotor_ratings,
        } for student in students]

        return render(request, self.template_name, {
            "form": form,
            "term": term,
            "classroom": classroom,
            "sheets": sheets,
        })


@method_decorator(it_staff_required, name="dispatch")
class PromotionResultView(View):
    """
    Prints promotion results (cumulative annual data) for each student in a class
    at the end of an academic session. Shows: cumulative score, average, position,
    and promotion status (Promoted/Promoted on Trial/Repeat).
    Optimized for printing with landscape mode and flexible scaling.
    """
    template_name = "results/staff/promotion_result.html"

    def get(self, request):
        session_id = request.GET.get("session")
        classroom_id = request.GET.get("classroom")

        form_kwargs = {}
        if session_id and classroom_id:
            form_kwargs["initial"] = {"session": session_id, "classroom": classroom_id}

        if not session_id or not classroom_id:
            return render(request, self.template_name, {"form": PromotionResultForm(**form_kwargs)})

        session = AcademicSession.objects.get(pk=session_id)
        classroom = ClassRoom.objects.get(pk=classroom_id)
        students = list(classroom.students.filter(is_active=True).order_by("last_name", "first_name"))

        # Gather promotion data for all students in the class
        promotions = []
        for student in students:
            session_result = SessionResult.objects.filter(
                student=student, session=session
            ).first()
            promotions.append({
                "student": student,
                "session_result": session_result,
            })

        return render(request, self.template_name, {
            "form": PromotionResultForm(**form_kwargs),
            "session": session,
            "classroom": classroom,
            "promotions": promotions,
        })
